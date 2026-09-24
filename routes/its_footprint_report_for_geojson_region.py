import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
import seaborn as sns
from sqlalchemy import text
from jinja2 import Template
import pdfkit
import base64
import os
from datetime import datetime
import io
from io import BytesIO
from dotenv.main import load_dotenv
import logging
import traceback
from fastapi import APIRouter, HTTPException
from fastapi.responses import HTMLResponse, StreamingResponse, Response
from sqlalchemy.orm import Session
from controller.db import SessionLocal, get_db
from routes.its_footprint_report_map_for_region import create_region_map
from fastapi import Depends
from pydantic import BaseModel
from typing import List
import json
import sys
import geopandas as gpd
from shapely.geometry import shape

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)
handler = logging.StreamHandler()
logger.addHandler(handler)

def debug_print(msg):
    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S.%f")[:-3]
    sys.stdout.write(f"[{timestamp}] {msg}\n")
    sys.stdout.flush()

router = APIRouter(tags=["Utility"], prefix='/Utility')

load_dotenv('fastApi/.env')

# Vegetation type to full name mapping
VEGETATION_MAPPING = {
    'AGRICULTURE': 'Agricultural Land',
    'FOREST': 'Forest Land',
    'GRASS_HERB': 'Grassland/Herbaceous',
    'SHRB_CHAP': 'Shrub/Chaparral',
    'SPARSE': 'Sparse Vegetation',
    'Trees Removed': 'Trees Removed',
    'URBAN': 'Urban/Developed',
    'WATER': 'Water',
    'WETLAND': 'Wetland'
}

# Ownership group to full name mapping
OWNERSHIP_MAPPING = {
    'FEDERAL': 'Federal Government',
    'LOCAL': 'Local Government',
    'NGO': 'Non-Governmental Organization',
    'PRIVATE_INDUSTRY': 'Private Industry',
    'PRIVATE_NON-INDUSTRY': 'Private Non-Industry',
    'STATE': 'State Government',
    'TRIBAL': 'Tribal Land'
}

class BackgroundLayer(BaseModel):
    layer_name: str
    layer_title: str

class GeoJSONRequest(BaseModel):
    geojson_str: str
    region_name: str
    output_format: str = 'html'
    background_layers: List[BackgroundLayer] = [
        {
            "layer_name": "rrk:annualburnprobability_202212_202406_t1_v5",
            "layer_title": "Annual Burn Probability"
        }
    ]

def get_footprint_data(db: Session, geojson_str: str, region_name: str):
    """Retrieve footprint data for a region defined by GeoJSON"""
    try:
        # Parse GeoJSON string to GeoDataFrame
        geojson_data = json.loads(geojson_str)
        geom = shape(geojson_data['features'][0]['geometry'])
        region_gdf = gpd.GeoDataFrame({'geometry': [geom]}, crs="EPSG:4326")
        debug_print(f"Parsed GeoJSON to GeoDataFrame: {region_gdf.shape}, CRS: {region_gdf.crs}")
        
        # Rename 'geometry' column to 'geom' for compatibility with create_region_map
        region_gdf = region_gdf.rename(columns={'geometry': 'geom'})
        # Set 'geom' as the active geometry column
        region_gdf = region_gdf.set_geometry('geom')
        
        # Transform to EPSG:3310 for querying
        region_gdf_3310 = region_gdf.to_crs("EPSG:3310")
        region_wkt = region_gdf_3310.geom.unary_union.wkt

        # Get footprint data intersecting with the GeoJSON geometry
        data_query = f"""
           WITH region_geom AS (
                SELECT ST_SetSRID(ST_GeomFromText('{region_wkt}'), 3310) AS geom
           ),
           bbox AS (
                SELECT ST_SetSRID(ST_Extent(geom), 3310)::geometry AS geom
                FROM region_geom
           )
           SELECT fp.*,
                  ST_X(ST_Transform(ST_Centroid(fp.geom), 4326)) as x,
                  ST_Y(ST_Transform(ST_Centroid(fp.geom), 4326)) as y
           FROM its.its_v2_0_footprint_points AS fp,
                bbox,
                region_geom
           WHERE ST_Within(fp.geom, bbox.geom)
             AND ST_Contains(region_geom.geom, fp.geom)
             AND year_txt ~ '^[0-9]+$'
             AND CAST(year_txt AS INTEGER) BETWEEN 2021 AND 2023
        """

        # print("-"*70)
        # print(data_query)

        with db.bind.connect() as conn:
            df = pd.read_sql_query(text(data_query), conn)
        debug_print(f"Footprint data retrieved: {df.shape}, columns: {list(df.columns)}")
        return df, region_gdf

    except Exception as e:
        logger.error(f"Error fetching footprint data: {str(e)}")
        debug_print(f"Error fetching footprint data: {str(e)}")
        return pd.DataFrame(), None

def plot_to_base64(plt_figure):
    buf = io.BytesIO()
    plt_figure.savefig(buf, format='png', dpi=300, bbox_inches='tight')
    buf.seek(0)
    img_str = base64.b64encode(buf.read()).decode('utf-8')
    buf.close()
    plt.close(plt_figure)
    return img_str

def adjust_chart_scaling(ax, data):
    max_val = data.max().max()
    min_val = data.min().min()
    if min_val == 0:
        min_val = 1e-6
    if max_val / min_val > 100:
        ax.set_yscale("log")
        ax.set_ylabel("Footprint Acres (Log Scale)")
    else:
        ax.set_ylabel("Footprint Acres")
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)

def create_ownership_footprint_chart(df):
    plt.style.use("seaborn-v0_8-whitegrid")
    df["ownership_name"] = df["primary_ownership_group"].map(OWNERSHIP_MAPPING).fillna(df["primary_ownership_group"])
    ownership_year_acres = df.groupby(["year_txt", "ownership_name"])['activity_quantity'].sum().unstack(fill_value=0)
    ownership_year_acres = ownership_year_acres.clip(upper=np.percentile(ownership_year_acres, 95))
    fig, ax = plt.subplots(figsize=(12, 6))
    ownership_year_acres.plot(kind="bar", ax=ax, width=0.8, colormap="coolwarm")
    ax.set_title("Footprint Acres by Land Ownership (2021-2023)", fontsize=14, fontweight="bold", pad=20)
    adjust_chart_scaling(ax, ownership_year_acres)
    ax.set_xlabel("Year", fontsize=12)
    ax.legend(title="Ownership", bbox_to_anchor=(1.05, 1), loc="upper left")
    plt.xticks(rotation=0)
    plt.tight_layout()
    return fig

def create_vegetation_footprint_chart(df):
    plt.style.use("seaborn-v0_8-whitegrid")
    df["vegetation_name"] = df["broad_vegetation_type"].map(VEGETATION_MAPPING).fillna(df["broad_vegetation_type"])
    veg_year_acres = df.groupby(["year_txt", "vegetation_name"])['activity_quantity'].sum().unstack(fill_value=0)
    veg_year_acres = veg_year_acres.clip(upper=np.percentile(veg_year_acres, 95))
    fig, ax = plt.subplots(figsize=(12, 6))
    veg_year_acres.plot(kind="bar", ax=ax, width=0.8, colormap="coolwarm")
    ax.set_title("Footprint Acres by Vegetation Types (2021-2023)", fontsize=14, fontweight="bold", pad=20)
    adjust_chart_scaling(ax, veg_year_acres)
    ax.set_xlabel("Year", fontsize=12)
    ax.legend(title="Vegetation Type", bbox_to_anchor=(1.05, 1), loc="upper left")
    plt.xticks(rotation=0)
    plt.tight_layout()
    return fig

def generate_footprint_summary_statistics(df, region_name):
    total_footprints = len(df)
    total_acres = df['activity_quantity'].sum()
    years_active = sorted(df[df['year_txt'].str.isnumeric()]['year_txt'].astype(int).unique())
    top_vegetation = df.groupby('broad_vegetation_type')['activity_quantity'].sum().idxmax()
    top_vegetation = VEGETATION_MAPPING.get(top_vegetation, top_vegetation)
    main_ownership = df.groupby('primary_ownership_group')['activity_quantity'].sum().idxmax()
    main_ownership = OWNERSHIP_MAPPING.get(main_ownership, main_ownership)
    return {
        'region_name': region_name,
        'total_footprints': total_footprints,
        'total_acres': f"{total_acres:,.0f}",
        'years_active': f"{min(years_active)} to {max(years_active)}" if years_active else "N/A",
        'top_vegetation': top_vegetation,
        'main_ownership': main_ownership
    }

def create_ownership_footprint_table(df):
    """Create styled table for ownership data by vegetation type"""
    ownership_df = df.groupby(
        ['primary_ownership_group', 'broad_vegetation_type', 'year_txt']
    )['activity_quantity'].sum().reset_index()

    ownership_df['Ownership'] = ownership_df['primary_ownership_group'].map(OWNERSHIP_MAPPING)
    ownership_df['Vegetation Type'] = ownership_df['broad_vegetation_type'].map(VEGETATION_MAPPING)
    ownership_df['Year'] = ownership_df['year_txt']
    ownership_df['Acres'] = ownership_df['activity_quantity'].apply(lambda x: f"{x:,.0f}")

    return ownership_df[['Ownership', 'Vegetation Type', 'Year', 'Acres']] \
        .sort_values(['Year', 'Ownership', 'Vegetation Type']) \
        .to_html(index=False, classes='table table-striped table-hover', border=0)

from openpyxl.utils import get_column_letter

def generate_footprint_excel_data(df, region_name):
    """Generate Excel data with a single sheet containing specified columns"""
    output = BytesIO()

    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        # Helper function for column width adjustment
        def auto_adjust_columns(worksheet):
            for column in worksheet.columns:
                max_length = 0
                column_letter = get_column_letter(column[0].column)
                for cell in column:
                    try:
                        cell_length = len(str(cell.value))
                        if cell_length > max_length:
                            max_length = cell_length
                    except:
                        pass
                adjusted_width = (max_length + 2) * 1.2
                worksheet.column_dimensions[column_letter].width = adjusted_width

        # Create single sheet with requested columns
        report_df = df.groupby(
            ['primary_ownership_group', 'broad_vegetation_type', 'year_txt']
        )['activity_quantity'].sum().reset_index()

        report_df['Ownership'] = report_df['primary_ownership_group'].map(OWNERSHIP_MAPPING)
        report_df['Vegetation Type'] = report_df['broad_vegetation_type'].map(VEGETATION_MAPPING)
        report_df['Year'] = report_df['year_txt'].astype(int)
        report_df['Acres'] = report_df['activity_quantity']

        report_df = report_df[[
            'Ownership',
            'Vegetation Type',
            'Year',
            'Acres'
        ]].sort_values(['Ownership', 'Vegetation Type', 'Year'])

        report_df.to_excel(writer, sheet_name='Footprint Report', index=False)
        auto_adjust_columns(writer.sheets['Footprint Report'])

        # Format numeric columns
        worksheet = writer.sheets['Footprint Report']
        for row in worksheet.iter_rows(min_row=2):
            for cell in row:
                if cell.column == 4:  # Acres column
                    cell.number_format = '#,##0.00'
                elif cell.column == 3:  # Year column
                    cell.number_format = '0'

        writer.book.active = 0

    output.seek(0)
    return output

@router.post("/its_footprint_report_for_geojson_region", include_in_schema=True)
def create_footprint_report_for_geojson_region(
    request: GeoJSONRequest,
    db: Session = Depends(get_db)
):
    geojson_str = request.geojson_str
    region_name = request.region_name
    output_format = request.output_format
    background_layers = request.background_layers

    debug_print(f"=======> background_layers: {background_layers}")
    debug_print(f"=======> region_name: {region_name}")
    debug_print(f"=======> output_format: {output_format}")
    debug_print(f"=======> geojson_str: {type(geojson_str)}")


    # Parse and validate GeoJSON
    try:
        geojson_data = json.loads(geojson_str)
        if not geojson_data.get('features') or not geojson_data['features'][0].get('geometry'):
            raise ValueError("Invalid GeoJSON: missing features or geometry")
        debug_print(f"GeoJSON parsed successfully, features: {len(geojson_data['features'])}")
    except json.JSONDecodeError as e:
        logger.error(f"Invalid GeoJSON format: {str(e)}")
        raise HTTPException(status_code=400, detail=f"Invalid GeoJSON format: {str(e)}")
    except ValueError as e:
        logger.error(f"GeoJSON validation error: {str(e)}")
        raise HTTPException(status_code=400, detail=f"GeoJSON validation error: {str(e)}")

    # Validate background layers
    try:
        debug_print(f"Received background_layers: {background_layers}")
        layers = [layer.dict() for layer in background_layers]
        debug_print(f"Parsed background_layers: {layers}")
        for layer in layers:
            BackgroundLayer(**layer)  # Validate each layer
    except Exception as e:
        logger.error(f"Invalid background_layers format: {str(e)}")
        raise HTTPException(status_code=400, detail=f"Invalid background_layers format: {str(e)}")

    # Get footprint data and region GeoDataFrame
    debug_print("Getting footprint points")
    df, region_gdf = get_footprint_data(db, geojson_str, region_name)
    if df.empty or region_gdf is None:
        logger.error(f"No footprint data found for region {region_name}")
        raise HTTPException(status_code=404,
                            detail=f"No footprint data found for region {region_name}")

    # Add Excel output handling
    if output_format.lower() == 'xlsx':
        debug_print("Generating Excel output")
        excel_data = generate_footprint_excel_data(df, region_name)
        return StreamingResponse(
            excel_data,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename={region_name}_footprint_report.xlsx"}
        )

    try:
        debug_print("Generating summary statistics and charts")
        ownership_chart = plot_to_base64(create_ownership_footprint_chart(df))
        vegetation_chart = plot_to_base64(create_vegetation_footprint_chart(df))
        summary = generate_footprint_summary_statistics(df, region_name)

        # Generate region maps for each background layer
        overlay_maps = []
        debug_print(f"Processing background layers: {len(layers)} layers found")
        if not layers:  # If background_layers is empty, generate default map
            debug_print("Background layers empty, generating default map for Footprint Locations")
            try:
                default_map = create_region_map(
                    db=db,
                    points_df=df,
                    table_name=None,
                    column_name=None,
                    region_name=region_name,
                    geoserver_url="https://sparcal.sdsc.edu/geoserver",
                    layer_name=None,
                    layer_title="Footprint Locations",
                    region_gdf=region_gdf
                )
                if default_map is None:
                    logger.error("Default map generation returned None")
                    debug_print("Default map generation returned None")
                    overlay_maps.append({
                        'image': None,
                        'title': 'Footprint Locations'
                    })
                else:
                    debug_print("Successfully generated default map for Footprint Locations")
                    overlay_maps.append({
                        'image': default_map,
                        'title': 'Footprint Locations'
                    })
            except Exception as e:
                logger.error(f"Default map generation error: {str(e)}")
                debug_print(f"Default map generation error: {str(e)}")
                traceback.print_exc()
                overlay_maps.append({
                    'image': None,
                    'title': 'Footprint Locations'
                })
        else:
            debug_print(f"Generating {len(layers)} overlay maps")
            for layer in layers:
                debug_print(f"Generating overlay_map for {layer['layer_title']}")
                try:
                    overlay_map = create_region_map(
                        db=db,
                        points_df=df,
                        table_name=None,
                        column_name=None,
                        region_name=region_name,
                        geoserver_url="https://sparcal.sdsc.edu/geoserver",
                        layer_name=layer['layer_name'],
                        layer_title=layer['layer_title'],
                        region_gdf=region_gdf
                    )
                    if overlay_map is None:
                        logger.error(f"Overlay map for {layer['layer_title']} returned None")
                        debug_print(f"Overlay map for {layer['layer_title']} returned None")
                        overlay_maps.append({
                            'image': None,
                            'title': layer['layer_title']
                        })
                    else:
                        debug_print(f"Generated overlay_map for {layer['layer_title']}")
                        overlay_maps.append({
                            'image': overlay_map,
                            'title': layer['layer_title']
                        })
                except Exception as e:
                    logger.error(f"Map generation error for {layer['layer_title']}: {str(e)}")
                    debug_print(f"Map generation error for {layer['layer_title']}: {str(e)}")
                    traceback.print_exc()
                    overlay_maps.append({
                        'image': None,
                        'title': layer['layer_title']
                    })

        debug_print(f"Generated overlay_maps: {len(overlay_maps)} maps")
        logger.info("Generating spreadsheet tables")
        ownership_table = create_ownership_footprint_table(df)

        # Generate HTML template
        debug_print("Rendering HTML template")
        html_template = Template('''
<!DOCTYPE html>
<html>
<head>
    <title>{{ summary.region_name }} Footprint Report</title>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <style>
        :root {
            --primary-color: #1e5c97;
            --secondary-color: #2c8b57;
            --accent-color: #f39c12;
            --light-bg: #f8f9fa;
            --dark-bg: #2c3e50;
            --text-color: #333333;
            --light-text: #ffffff;
            --border-color: #e0e0e0;
            --box-shadow: 0 4px 6px rgba(0, 0, 0, 0.1);
        }
        body {
            font-family: 'Segoe UI', Roboto, 'Helvetica Neue', Arial, sans-serif;
            line-height: 1.6;
            color: var(--text-color);
            max-width: 1200px;
            margin: 0 auto;
            padding: 0;
            background-color: #f5f7fa;
        }
        h1, h2, h3, h4 {
            font-weight: 600;
            color: var(--primary-color);
            margin-top: 1.5em;
            margin-bottom: 0.8em;
        }
        h1 {
            font-size: 2.5rem;
            margin-top: 0;
        }
        h2 {
            font-size: 1.8rem;
            border-bottom: 2px solid var(--border-color);
            padding-bottom: 0.3em;
        }
        h3 {
            font-size: 1.4rem;
            color: var(--secondary-color);
        }
        .container {
            background-color: white;
            box-shadow: var(--box-shadow);
            padding: 2rem;
            margin: 0 auto;
        }
        .header {
            background: linear-gradient(135deg, var(--primary-color), var(--secondary-color));
            color: var(--light-text);
            padding: 3rem;
            margin-bottom: 2rem;
            text-align: center;
            position: relative;
            box-shadow: 0 4px 12px rgba(0, 0, 0, 0.15);
        }
        .header h1, .header h2 {
            color: white;
            margin: 0.5rem 0;
            text-shadow: 1px 1px 3px rgba(0, 0, 0, 0.2);
        }
        .header p {
            margin-top: 1rem;
            font-size: 1.1rem;
            opacity: 0.9;
        }
        .section {
            margin-bottom: 2.5rem;
            padding: 0 1.5rem;
        }
        .introduction {
            background-color: var(--light-bg);
            border-left: 4px solid var(--primary-color);
            padding: 1.5rem;
            margin-bottom: 2rem;
            border-radius: 0 4px 4px 0;
        }
        .summary-box {
            background-color: white;
            border-radius: 8px;
            box-shadow: var(--box-shadow);
            padding: 1.5rem 2rem;
            margin-bottom: 2rem;
            border-top: 5px solid var(--accent-color);
        }
        .summary-box h2 {
            color: var(--accent-color);
            border-bottom: none;
            margin-top: 0;
        }
        .summary-stats {
            display: flex;
            flex-wrap: wrap;
            gap: 1.5rem;
            margin-top: 1.5rem;
        }
        .stat-item {
            flex: 1;
            min-width: 200px;
            background-color: var(--light-bg);
            padding: 1rem;
            border-radius: 6px;
            text-align: center;
        }
        .stat-value {
            font-size: 1.8rem;
            fontweight: 700;
            color: var(--primary-color);
            margin-bottom: 0.5rem;
        }
        .stat-label {
            font-size: 0.9rem;
            color: #666;
            text-transform: uppercase;
            letter-spacing: 0.5px;
        }
        .chart-container {
            margin: 2rem 0;
        }
        .chart-row {
            display: flex;
            flex-wrap: wrap;
            gap: 2rem;
            margin-bottom: 2rem;
        }
        .chart {
            flex: 1;
            min-width: 300px;
            background-color: white;
            border-radius: 8px;
            box-shadow: var(--box-shadow);
            padding: 1.5rem;
            transition: transform 0.2s ease;
        }
        .chart:hover {
            transform: translateY(-5px);
        }
        .chart h3 {
            text-align: center;
            margin-top: 0;
            padding-bottom: 0.8rem;
            border-bottom: 1px solid var(--border-color);
        }
        .chart img {
            max-width: 100%;
            height: auto;
            display: block;
            margin: 1rem auto;
        }
        .chart-caption {
            font-size: 0.9rem;
            color: #666;
            text-align: center;
            margin-top: 1rem;
            font-style: italic;
        }
        .full-width {
            flex-basis: 100%;
        }
        table {
            border-collapse: collapse;
            width: 100%;
            margin: 2rem 0;
            background-color: white;
            box-shadow: var(--box-shadow);
            border-radius: 8px;
            overflow: hidden;
        }
        th, td {
            padding: 12px 15px;
            text-align: left;
        }
        th {
            background-color: var(--primary-color);
            color: white;
            font-weight: 600;
            text-transform: uppercase;
            font-size: 0.9rem;
            letter-spacing: 0.5px;
        }
        tr:nth-child(even) {
            background-color: #f2f7ff;
        }
        tr:hover {
            background-color: #e6f0ff;
        }
        .findings-box {
            background-color: var(--light-bg);
            padding: 1.5rem 2rem;
            border-radius: 8px;
            margin-bottom: 2rem;
            border-left: 4px solid var(--primary-color);
        }
        .recommendations {
            background-color: var(--light-bg);
            padding: 1.5rem 2rem;
            border-radius: 8px;
            border-left: 4px solid var(--secondary-color);
        }
        .recommendations ul {
            padding-left: 1.2rem;
        }
        .recommendations li {
            margin-bottom: 0.8rem;
        }
        .footer {
            margin-top: 3rem;
            padding: 2rem 0;
            text-align: center;
            font-size: 0.9rem;
            color: #777;
            border-top: 1px solid var(--border-color);
        }
        @media print {
            body {
                background-color: white;
            }
            .container {
                box-shadow: none;
                padding: 0;
            }
            .chart:hover {
                transform: none;
            }
            .chart-row {
                display: block;
            }
            .chart {
                width: 100%;
                margin-bottom: 2rem;
                box-shadow: none;
                page-break-inside: avoid;
            }
            .header {
                background: var(--primary-color) !important;
                -webkit-print-color-adjust: exact;
            }
            th {
                background-color: var(--primary-color) !important;
                color: white !important;
                -webkit-print-color-adjust: exact;
            }
            .summary-box {
                box-shadow: none;
                border: 1px solid var(--border-color);
            }
            table {
                box-shadow: none;
            }
            /* Chart sizing adjustments */
            .chart img {
                max-height: 320px !important;
                width: auto !important;
                margin: 12px auto !important;
                page-break-inside: avoid;
            }
            /* Container adjustments */
            .chart {
                page-break-inside: avoid;
                margin: 8px 0 !important;
                padding: 4px !important;
            }
            /* Grid layout for charts */
            .chart-row {
                display: grid !important;
                grid-template-columns: 1fr 1fr;
                gap: 8px !important;
                page-break-inside: avoid;
            }
            /* Single chart full width */
            .full-width.chart img {
                max-height: 400px !important;
                width: 95% !important;
            }
        }
        /* Chart image quality preservation */
        .chart img {
            image-rendering: crisp-edges;
            -webkit-print-color-adjust: exact;
            print-color-adjust: exact;
        }
        /* PDF-specific image scaling */
        @media print {
            canvas {
                max-width: 100% !important;
                height: auto !important;
            }
            figure {
                max-width: 90% !important;
                margin: 0 auto !important;
            }
        }
    </style>
</head>
<body>
    <div class="container">
        <div class="header">
            <p>Wildfire & Forest Resilience Task Force</p>
            <h2>Vegetation Treatment Footprint Report</h2>
            <p>
                <span style="font-size:16pt; font-weight: bold;">{{ summary.region_name }}</span>
                <br/>
                Custom GeoJSON Region
            </p>
        </div>

        <div class="section">
            <div class="summary-box">
                <h2>Region Footprint Overview</h2>
                <div class="summary-stats">
                    <div class="stat-item">
                        <div class="stat-value">{{ summary.total_footprints }}</div>
                        <div class="stat-label">Total Footprint Entries</div>
                    </div>
                    <div class="stat-item">
                        <div class="stat-value">{{ summary.total_acres }}</div>
                        <div class="stat-label">Total Footprint Acres</div>
                    </div>
                    <div class="stat-item">
                        <div class="stat-value">{{ summary.years_active }}</div>
                        <div class="stat-label">Active Period</div>
                    </div>
                </div>
            </div>
        </div>

        <div class="section">
            <div class="chart-container">
                {% for map in overlay_maps %}
                <div class="chart-row">
                    <div class="chart full-width">
                        <h3>Footprint Locations{% if map.title != 'Footprint Locations' %} and {{ map.title }}{% endif %}</h3>
                        {% if map.image %}
                            <img src="data:image/png;base64,{{ map.image }}" alt="{{ map.title }} Overlay">
                        {% else %}
                            <p>Overlay map for {{ map.title }} could not be generated.</p>
                        {% endif %}
                        {% if map.title != "Footprint Locations" %}                                                                                                     
                            <p class="table-description" style="text-align: center;">                                                                                   
                                {{ map.title }} - <a href="https://caregionalresourcekits.org/clm.html" target="_blank">California Landscape Metrics</a>                
                            </p>                                                                                                                                        
                        {% endif %}       
                    </div>
                </div>
                {% endfor %}
            </div>
        </div>

        {% if overlay_maps|length % 2 == 1 %}
        <div class="section" style="margin-top: 5cm;">
        {% else %}
        <div class="section">
        {% endif %}
            <h2>Treatment Footprint Analysis</h2>
            <div class="chart-container">
                <div class="chart-row">
                    <div class="chart">
                        <h3>Land Ownership Distribution</h3>
                        <img src="data:image/png;base64,{{ ownership_chart }}" alt="Footprint by Ownership">
                    </div>
                </div>
                <div class="chart-row">
                    <div class="chart">
                        <h3>Vegetation Types Treated</h3>
                        <img src="data:image/png;base64,{{ vegetation_chart }}" alt="Vegetation Types Treated">
                    </div>
                </div>
            </div>
        </div>

        <div class="section">
            <h2>Detailed Footprint Data</h2>
            <h3>Land Ownership Details</h3>
            <p class="table-description">
                Distribution of footprints across land ownership types, showing annual implementation rates by vegetation types.
            </p>
            {{ ownership_table|safe }}
        </div>

        <div class="footer">
            <p><strong>California Wildfire and Forest Interagency Treatment Dashboard</strong></p>
            <p>Report Generated: {{ current_date }}</p>
        </div>
    </div>
</body>
</html>
        ''')

        # Render HTML
        html_content = html_template.render(
            summary=summary,
            overlay_maps=overlay_maps,
            ownership_chart=ownership_chart,
            vegetation_chart=vegetation_chart,
            ownership_table=ownership_table,
            current_date=datetime.now().strftime('%B %d, %Y')
        )

        # Handle output format
        debug_print(f"Returning output format: {output_format.lower()}")
        if output_format.lower() == 'html':
            return HTMLResponse(content=html_content)
        elif output_format.lower() == 'pdf':
            options = {
                'page-size': 'Letter',
                'margin-top': '0.5in',
                'margin-right': '0.5in',
                'margin-bottom': '0.5in',
                'margin-left': '0.5in',
                'encoding': "UTF-8",
                'enable-local-file-access': None
            }
            pdf = pdfkit.from_string(html_content, False, options=options)
            return Response(content=pdf, media_type="application/pdf")

        logger.error(f"Invalid output format: {output_format}")
        raise HTTPException(status_code=400, detail="Invalid output format")

    except Exception as e:
        logger.error(f"Report generation failed: {str(e)}")
        debug_print(f"Report generation failed: {str(e)}")
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=str(e))

